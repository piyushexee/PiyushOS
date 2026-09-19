"""Tools ka quick test - bina server/phone ke.

    .venv/bin/python -m brain.dev.test_tools
"""
import json
import os
import sys

from ..tools.edit_tool import edit_presentation, edit_spreadsheet
from ..tools.model3d_tool import build_model3d
from ..tools.pptx_tool import build_pptx
from ..tools.xlsx_tool import build_xlsx
from ..registry import LAST_FILES

OUTLINE = json.dumps([
    {"title": "Test PPT", "subtitle": "PiyushOS demo"},
    {"title": "Point 1", "bullets": ["Alpha", "Beta", "Gamma"]},
    {"title": "Point 2", "bullets": ["X", "Y"]},
])

SHEET = json.dumps({
    "title": "Test Sheet",
    "sheets": [{"name": "Data", "headers": ["A", "B", "C"],
                "rows": [[1, 2, 3], ["chotu", "bada", 4]]}],
})

PASS = 0
FAIL = 0


def check(label: str, cond: bool, extra: str = ""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"✅ {label} {extra}")
    else:
        FAIL += 1
        print(f"❌ {label} {extra}")


def main():
    # ---------- 1. PPTX ----------
    res = build_pptx("AI", OUTLINE)
    check("pptx create", os.path.exists(res["path"]), res["name"])

    # ---------- 2. XLSX ----------
    res = build_xlsx(SHEET)
    check("xlsx create", os.path.exists(res["path"]), res["name"])

    # ---------- 3. 3D models ----------
    for shape in ("rocket", "house", "car", "chair", "combo", "torus"):
        try:
            r = build_model3d("m", shape, "red")
            ok = all(os.path.exists(f) and os.path.getsize(f) > 500 for f in r["files"])
            check(f"3d {shape}", ok, r["name"])
        except Exception as e:
            check(f"3d {shape}", False, str(e))

    # ---------- 4. PPT EDIT ----------
    try:
        e = edit_presentation(json.dumps({"file": "auto", "edits": [
            {"op": "set_title", "slide": 2, "text": "NAYA Title"},
            {"op": "add_slide", "after": 2, "title": "Extra Slide", "bullets": ["Zeta"]},
            {"op": "delete_slide", "slide": 3},
        ]}))
        from pptx import Presentation
        prs = Presentation(e["path"])
        n = len(prs.slides._sldIdLst)
        check("pptx edit", os.path.exists(e["path"]) and n == 4, f"({n} slides)")
    except Exception as ex:
        check("pptx edit", False, str(ex))

    # ---------- 5. XLSX EDIT ----------
    try:
        e = edit_spreadsheet(json.dumps({"file": "auto", "edits": [
            {"op": "set_cell", "sheet": "Data", "cell": "B3", "value": 99},
            {"op": "add_row", "sheet": "Data", "row": ["naya", 5, 6]},
        ]}))
        from openpyxl import load_workbook
        wb = load_workbook(e["path"])
        val = wb["Data"]["B3"].value
        check("xlsx edit", os.path.exists(e["path"]) and val == 99, f"(B3={val})")
    except Exception as ex:
        check("xlsx edit", False, str(ex))

    # ---------- 6. Registry ----------
    check("last files registry", "pptx" in LAST_FILES and "xlsx" in LAST_FILES)

    print(f"\n{PASS}/{PASS + FAIL} checks pass")
    sys.exit(0 if FAIL == 0 else 1)


if __name__ == "__main__":
    main()
