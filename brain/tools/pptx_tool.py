"""PPTX generator (python-pptx) - clean dark theme ke saath."""
import json
import time

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from ..config import settings
from ..registry import LAST_FILES

NAVY = RGBColor(0x14, 0x1E, 0x3C)
ACCENT = RGBColor(0x4F, 0xC3, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x20, 0x2A, 0x44)
GREY = RGBColor(0x8A, 0x97, 0xB5)


def _set_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _add_rect(slide, x, y, w, h, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _safe_name(s: str, limit: int = 40) -> str:
    s = "".join(ch for ch in (s or "") if ch.isalnum() or ch in " -_").strip()
    return s.replace(" ", "_")[:limit] or "file"


def add_content_slide(prs, blank, title, bullets):
    """Ek content slide banata hai (editing tools isi ko reuse karte hain)."""
    s = prs.slides.add_slide(blank)
    _set_bg(s, WHITE)
    _add_rect(s, 0, 0, prs.slide_width, Inches(1.35), NAVY)
    tb = s.shapes.add_textbox(Inches(0.9), Inches(0.33), Inches(11.5), Inches(0.8))
    rt = tb.text_frame.paragraphs[0].add_run()
    rt.text = title
    rt.font.size = Pt(30)
    rt.font.bold = True
    rt.font.color.rgb = WHITE
    _add_rect(s, Inches(0.9), Inches(1.35), Inches(1.6), Inches(0.07), ACCENT)
    body = s.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(11.5), Inches(5.1))
    tf = body.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets[:8]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        r = p.add_run()
        r.text = f"▪  {b}"
        r.font.size = Pt(22)
        r.font.color.rgb = DARK
    return s


def build_pptx(topic: str, outline_json: str) -> dict:
    """Banata hai PPT aur {path, files, summary} return karta hai."""
    try:
        outline = json.loads(outline_json) if isinstance(outline_json, str) else outline_json
        if isinstance(outline, dict):
            outline = [outline]
        if not outline:
            raise ValueError("outline khaali hai")
    except Exception as e:
        raise ValueError(f"outline_json valid JSON array nahi hai: {e}")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    count = 0

    # ---------- Title slide ----------
    t = outline[0]
    s = prs.slides.add_slide(blank)
    _set_bg(s, NAVY)
    _add_rect(s, Inches(0.9), Inches(4.4), Inches(2.2), Inches(0.09), ACCENT)
    box = s.shapes.add_textbox(Inches(0.9), Inches(2.5), Inches(11.5), Inches(1.8))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = t.get("title", topic)
    r.font.size = Pt(50)
    r.font.bold = True
    r.font.color.rgb = WHITE
    if t.get("subtitle"):
        box2 = s.shapes.add_textbox(Inches(0.9), Inches(4.7), Inches(11.5), Inches(0.8))
        p2 = box2.text_frame.paragraphs[0]
        r2 = p2.add_run()
        r2.text = t["subtitle"]
        r2.font.size = Pt(22)
        r2.font.color.rgb = ACCENT
    foot = s.shapes.add_textbox(Inches(0.9), Inches(6.75), Inches(11.5), Inches(0.5))
    rf = foot.text_frame.paragraphs[0].add_run()
    rf.text = "PiyushOS se banaya hua"
    rf.font.size = Pt(12)
    rf.font.color.rgb = GREY
    count += 1

    # ---------- Content slides ----------
    for item in outline[1:]:
        title = item.get("title", "")
        bullets = item.get("bullets", [])
        if isinstance(bullets, str):
            bullets = [bullets]
        add_content_slide(prs, blank, title, bullets)
        count += 1

    # ---------- Thank you slide ----------
    s = prs.slides.add_slide(blank)
    _set_bg(s, NAVY)
    tb = s.shapes.add_textbox(Inches(1.5), Inches(3.0), Inches(10.3), Inches(1.5))
    para = tb.text_frame.paragraphs[0]
    rt = para.add_run()
    rt.text = "Shukriya!"
    rt.font.size = Pt(48)
    rt.font.bold = True
    rt.font.color.rgb = WHITE
    para.alignment = PP_ALIGN.CENTER
    count += 1

    stamp = time.strftime("%Y%m%d_%H%M%S")
    path = settings.out_dir / f"presentation_{_safe_name(topic) or 'slides'}_{stamp}.pptx"
    prs.save(path)
    LAST_FILES["pptx"] = str(path)
    return {
        "path": str(path),
        "name": path.name,
        "slides": count,
        "files": [str(path)],
        "summary": f"PPT ban gayi: {path.name} ({count} slides)",
    }
