"""XLSX generator (openpyxl) - styled headers ke saath."""
import json
import time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..config import settings
from ..registry import LAST_FILES

HDR_FONT = Font(bold=True, color="FFFFFF", size=12)
NAVY_FILL = PatternFill("solid", fgColor="141E3C")


def _safe_name(s: str, limit: int = 40) -> str:
    s = "".join(ch for ch in (s or "") if ch.isalnum() or ch in " -_").strip()
    return s.replace(" ", "_")[:limit] or "sheet"


def build_xlsx(spec_json: str) -> dict:
    try:
        spec = json.loads(spec_json) if isinstance(spec_json, str) else spec_json
        sheets = spec.get("sheets", [])
        if not sheets:
            raise ValueError("'sheets' list khaali hai")
    except Exception as e:
        raise ValueError(f"spec_json valid nahi hai: {e}")

    wb = Workbook()
    wb.remove(wb.active)

    for sh in sheets[:8]:
        name = str(sh.get("name", "Sheet"))[:31] or "Sheet"
        headers = sh.get("headers", []) or []
        rows = sh.get("rows", []) or []
        ws = wb.create_sheet(name)

        if headers:
            ws.append(headers)
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=c)
                cell.font = HDR_FONT
                cell.fill = NAVY_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in rows[:500]:
            ws.append(list(row))

        # column widths
        for c in range(1, len(headers) + 1):
            vals = [len(str(headers[c - 1]))]
            vals += [len(str(r[c - 1])) for r in rows[:200] if c - 1 < len(r)]
            ws.column_dimensions[get_column_letter(c)].width = min(max(max(vals) + 3, 10), 42)

        if headers:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 2)}"
            ws.freeze_panes = "A2"

    stamp = time.strftime("%Y%m%d_%H%M%S")
    title = str(spec.get("title", "spreadsheet"))
    path = settings.out_dir / f"{_safe_name(title)}_{stamp}.xlsx"
    wb.save(path)
    LAST_FILES["xlsx"] = str(path)
    return {
        "path": str(path),
        "name": path.name,
        "sheets": [str(s.get("name", "Sheet")) for s in sheets],
        "files": [str(path)],
        "summary": f"Excel ban gaya: {path.name} (sheets: {', '.join(str(s.get('name')) for s in sheets)})",
    }
